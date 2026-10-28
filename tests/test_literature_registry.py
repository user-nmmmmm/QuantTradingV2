from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from analysis.literature_registry import (
    EXPECTED_EXCEL_IDS, PRIMARY_SOURCES, build_registry, normalize_title, presence_differs_from_expected,
    read_workbook_rows, validate_registry,
)
from scripts.build_paper_literature_registry import main

ROOT = Path(__file__).resolve().parents[1]
COMMITTED = ROOT / "docs/research/paper_literature_registry_20261026.json"
HEADER = ["编号", "大分类", "小分类", "论文题目", "作者", "年份", "被引次数", "阅读参考", "出版信息", "研究切入点",
          "摘要（检索原文）", "主检索方向", "全部检索方向", "论文链接", "PDF链接", "Source Name", "Source URL"]


def build_small(path):
    return build_registry(*read_workbook_rows(path), supplements=(), reserve_candidates=())


def make_workbook(path, titles, ids=None):
    wb = openpyxl.Workbook()
    wb.active.title = "使用说明"
    wb.active["A2"] = "基于 Scholar 学术检索整理｜检索日期：2026-10-25｜共 1 个大类"
    sheet = wb.create_sheet("论文总表")
    sheet.append(["论文总表"])
    sheet.append(["说明"])
    sheet.append([])
    sheet.append(HEADER)
    for index, title in enumerate(titles):
        number = (ids or list(range(1, len(titles) + 1)))[index]
        sheet.append([number, "大类", "小类", title, "A Author", 2020, 10, "补充阅读", "pub", "hook", "snippet", "x", "x",
                      f"https://example.org/{number}", "https://example.org/p.pdf" if number % 2 else None, "S", "u"])
    wb.save(path)


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "papers.xlsx"
    titles = ["The deflated Sharpe ratio: Correcting for selection bias", "Optimal execution of portfolio transactions",
              "Something unrelated"]
    make_workbook(path, titles, ids=[12, 252, 300])
    return path


def test_reads_rows_and_hashes_the_workbook_bytes(workbook):
    rows, source = read_workbook_rows(workbook)
    assert [r["excel_id"] for r in rows] == ["12", "252", "300"]
    assert source["row_count"] == 3 and source["retrieved_on"] == "2026-10-25"
    import hashlib
    assert source["sha256"] == hashlib.sha256(workbook.read_bytes()).hexdigest()


def test_presence_is_computed_from_titles_not_assumed(workbook):
    registry = build_small(workbook)
    present = {p["id"]: p["excel_ids"] for p in registry["primary_sources"]}
    assert present["L02"] == [12] and present["L04"] == [252]
    assert all(not present[f"L0{i}"] for i in (1, 3, 5, 6, 7, 8))
    assert presence_differs_from_expected(registry) == []
    # a workbook where the DSR paper is absent is reported, not silently accepted
    other = workbook.with_name("other.xlsx")
    make_workbook(other, ["Optimal execution of portfolio transactions"], ids=[252])
    assert presence_differs_from_expected(build_small(other)) == ["L02: found [], expected 12"]


def test_pdf_flag_and_reading_level_do_not_overstate(workbook):
    registry = build_small(workbook)
    first = registry["papers"][0]
    assert first["has_pdf_link"] is False and first["reading_level"] == "metadata_and_search_snippet_only"
    assert all(p["reading_level"] == "source_located_full_text_not_obtained" for p in registry["primary_sources"])
    assert registry["claims"]["full_excel_mapping_of_primary_sources"] is False


def test_missing_sheet_or_columns_are_rejected(tmp_path):
    path = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "other"
    wb.save(path)
    with pytest.raises(ValueError, match="not found"):
        read_workbook_rows(path)
    wb = openpyxl.Workbook()
    wb.active.title = "论文总表"
    for _ in range(3):
        wb.active.append([])
    wb.active.append(["编号", "论文题目"])
    wb.save(path)
    with pytest.raises(ValueError, match="missing columns"):
        read_workbook_rows(path)


def test_build_refuses_unknown_supplement_ids(workbook):
    # the real supplements reference ids 1..288; a tiny workbook cannot satisfy them
    with pytest.raises(ValueError, match="unknown ids"):
        build_registry(*read_workbook_rows(workbook))


def test_committed_registry_is_internally_consistent():
    registry = json.loads(COMMITTED.read_text(encoding="utf-8"))
    assert validate_registry(registry) == []
    assert presence_differs_from_expected(registry) == []
    assert len(registry["papers"]) == registry["source_workbook"]["row_count"] == 288
    by_id = {p["excel_id"]: p for p in registry["papers"]}
    for excel_id in EXPECTED_EXCEL_IDS.values():
        assert by_id[excel_id]["has_pdf_link"] is True
    assert [p["id"] for p in registry["primary_sources"]] == [item[0] for item in PRIMARY_SOURCES]


def test_validation_catches_tampering():
    registry = json.loads(COMMITTED.read_text(encoding="utf-8"))
    registry["primary_sources"][0]["in_excel"] = True
    assert any("in_excel disagrees" in p for p in validate_registry(registry))
    registry = json.loads(COMMITTED.read_text(encoding="utf-8"))
    registry["primary_sources"][1]["excel_ids"] = [13]
    assert any("does not match the title" in p for p in validate_registry(registry))
    registry = json.loads(COMMITTED.read_text(encoding="utf-8"))
    registry["papers"].append(dict(registry["papers"][0]))
    problems = validate_registry(registry)
    assert any("duplicate" in p for p in problems) and any("row_count" in p for p in problems)
    registry = json.loads(COMMITTED.read_text(encoding="utf-8"))
    registry["supplements"][0]["excel_ids"].append(9999)
    assert any("unknown ids" in p for p in validate_registry(registry))


def test_cli_leaves_existing_output_untouched_when_build_fails(tmp_path, workbook, capsys):
    out = tmp_path / "out.json"
    out.write_text("{}", encoding="utf-8")
    # tiny workbook cannot satisfy the supplement references, so build fails before writing either way
    with pytest.raises(ValueError):
        main(["build", "--workbook", str(workbook), "--output", str(out)])
    assert out.read_text(encoding="utf-8") == "{}"


def test_normalize_title_ignores_case_and_punctuation():
    assert normalize_title("Honey, I Shrunk the Sample Covariance Matrix") == "honey i shrunk the sample covariance matrix"
