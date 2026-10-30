"""Installed runtime dependencies preserve real client, signing and report behavior."""
from __future__ import annotations

import asyncio
import base64
import gzip
import hashlib
import hmac
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import importlib
import importlib.metadata
import io
import re
import ssl
from pathlib import Path
import sys
from threading import Thread
import zlib

from packaging.requirements import Requirement
from packaging.utils import canonicalize_name
from packaging.version import Version
import pytest
import requests


ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("distribution,minimum", [
    ("aiohttp", "3.14.4"), ("certifi", "2026.10.19"),
    ("cryptography", "50.0.2"), ("Pillow", "13.0.0"),
    ("soupsieve", "3.0.3"), ("urllib3", "2.8.0"), ("requests", "2.33.1"),
    ("idna", "3.15"), ("multidict", "6.9.1"),
])
def test_installed_client_and_report_versions_meet_the_runtime_floor(distribution, minimum):
    assert Version(importlib.metadata.version(distribution)) >= Version(minimum)


def test_requests_uses_the_installed_certificate_bundle():
    import certifi

    bundle = Path(certifi.where())
    assert Path(requests.certs.where()).resolve() == bundle.resolve()
    certificates = re.findall(
        rb"-----BEGIN CERTIFICATE-----.*?-----END CERTIFICATE-----", bundle.read_bytes(), re.S)
    assert certificates
    fingerprints = {
        hashlib.sha256(ssl.PEM_cert_to_DER_cert(block.decode("ascii"))).hexdigest()
        for block in certificates
    }
    assert "86a1ecba089c4a8d3bbe2734c612ba341d813e043cf9e8a862cd5c57a36bbe6b" not in fingerprints
    with requests.Session() as session:
        assert session.verify is True


def _localhost_tls_material(tmp_path):
    from datetime import datetime, timedelta, timezone
    from ipaddress import ip_address

    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.x509.oid import ExtendedKeyUsageOID, NameOID

    now = datetime.now(timezone.utc)
    ca_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    ca_name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "Local test CA")])
    ca = (x509.CertificateBuilder().subject_name(ca_name).issuer_name(ca_name)
          .public_key(ca_key.public_key()).serial_number(x509.random_serial_number())
          .not_valid_before(now - timedelta(minutes=5)).not_valid_after(now + timedelta(days=1))
          .add_extension(x509.BasicConstraints(ca=True, path_length=0), critical=True)
          .add_extension(x509.KeyUsage(False, False, False, False, False, True, True, False, False),
                         critical=True)
          .add_extension(x509.SubjectKeyIdentifier.from_public_key(ca_key.public_key()), critical=False)
          .sign(ca_key, hashes.SHA256()))
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    leaf = (x509.CertificateBuilder()
            .subject_name(x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "localhost")]))
            .issuer_name(ca_name).public_key(key.public_key())
            .serial_number(x509.random_serial_number())
            .not_valid_before(now - timedelta(minutes=5)).not_valid_after(now + timedelta(days=1))
            .add_extension(x509.BasicConstraints(ca=False, path_length=None), critical=True)
            .add_extension(x509.SubjectAlternativeName([x509.IPAddress(ip_address("127.0.0.1"))]),
                           critical=False)
            .add_extension(x509.ExtendedKeyUsage([ExtendedKeyUsageOID.SERVER_AUTH]), critical=False)
            .add_extension(x509.AuthorityKeyIdentifier.from_issuer_public_key(ca_key.public_key()),
                           critical=False)
            .add_extension(x509.KeyUsage(True, False, True, False, False, False, False, False, False),
                           critical=True)
            .sign(ca_key, hashes.SHA256()))
    ca_path, cert_path, key_path = (tmp_path / name for name in ("ca.pem", "server.pem", "key.pem"))
    ca_path.write_bytes(ca.public_bytes(serialization.Encoding.PEM))
    cert_path.write_bytes(leaf.public_bytes(serialization.Encoding.PEM))
    key_path.write_bytes(key.private_bytes(serialization.Encoding.PEM,
        serialization.PrivateFormat.PKCS8, serialization.NoEncryption()))
    return ca_path, cert_path, key_path


def test_requests_rejects_untrusted_tls_and_accepts_explicit_local_ca(tmp_path):
    ca_path, cert_path, key_path = _localhost_tls_material(tmp_path)

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"verified-local-response")

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    context.load_cert_chain(certfile=cert_path, keyfile=key_path)
    server.socket = context.wrap_socket(server.socket, server_side=True)
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    try:
        with requests.Session() as session:
            session.trust_env = False
            address = f"https://127.0.0.1:{server.server_port}/"
            with pytest.raises(requests.exceptions.SSLError):
                session.get(address, timeout=2)
            assert session.get(address, verify=str(ca_path), timeout=2).text == "verified-local-response"
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)
    assert not worker.is_alive()


@pytest.mark.parametrize("ca_variable", [None, "REQUESTS_CA_BUNDLE", "CURL_CA_BUNDLE"])
def test_public_archive_preserves_explicit_certificate_trust(tmp_path, monkeypatch, ca_variable):
    from scripts import collect_factor_evidence as collector

    ca_path, cert_path, key_path = _localhost_tls_material(tmp_path)

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"verified-public-source")

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    context.load_cert_chain(certfile=cert_path, keyfile=key_path)
    server.socket = context.wrap_socket(server.socket, server_side=True)
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    address = f"https://127.0.0.1:{server.server_port}/"
    monkeypatch.setattr(collector, "BASE", address)
    monkeypatch.setenv("NO_PROXY", "127.0.0.1")
    monkeypatch.setenv("no_proxy", "127.0.0.1")
    for name in ("REQUESTS_CA_BUNDLE", "CURL_CA_BUNDLE"):
        monkeypatch.delenv(name, raising=False)
    if ca_variable is not None:
        monkeypatch.setenv(ca_variable, str(ca_path))
    try:
        with requests.Session() as session:
            archive = collector.PublicArchive(tmp_path / "archive", session=session, timeout=2)
            raw, metadata = archive.fetch("tls-source", address)
            if ca_variable is None:
                assert raw == b"" and metadata["error_type"] == "SSLError"
                assert metadata["status"] is None
            else:
                assert raw == b"verified-public-source"
                assert metadata["status"] == 200 and metadata["error_type"] is None
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)
    assert not worker.is_alive()


def test_editable_install_metadata_and_lock_preserve_the_supported_dependency_boundary():
    declared = {canonicalize_name(line.split("==")[0]) for line in
                (ROOT / "requirements.txt").read_text().splitlines() if "==" in line}
    installed = {canonicalize_name(requirement.name) for text in
                 importlib.metadata.requires("quant-trading") or []
                 if (requirement := Requirement(text)).marker is None
                 or requirement.marker.evaluate({"extra": ""})}
    assert installed == declared
    locked = {canonicalize_name(line.split("==")[0]) for line in
              (ROOT / "requirements.lock.txt").read_text().splitlines() if "==" in line}
    assert not {"streamlit", "gitpython", "gitdb", "smmap", "tornado", "pyarrow"} & locked


def test_direct_environment_check_uses_distribution_import_aliases(monkeypatch, capsys):
    from scripts import check_environment

    monkeypatch.setattr(sys, "argv", ["check_environment"])
    assert not check_environment.main()
    assert "Environment OK" in capsys.readouterr().out


@pytest.mark.parametrize("version,rejected", [((3, 11, 4), True), ((3, 11, 5), False), ((3, 13, 2), False)])
def test_environment_enforces_the_parser_python_patch_floor(monkeypatch, capsys, version, rejected):
    import tomllib

    from scripts import check_environment

    metadata = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))
    assert metadata["project"]["requires-python"] == ">=3.11.5"
    monkeypatch.setattr(sys, "argv", ["check_environment"])
    monkeypatch.setattr(sys, "version_info", version)
    monkeypatch.setattr(check_environment, "pins", lambda path: [])
    assert bool(check_environment.main()) is rejected
    assert ("Python 3.11.5+ is required" in capsys.readouterr().err) is rejected


@pytest.mark.parametrize("encoding", ["gzip", "deflate"])
def test_requests_decodes_public_archive_stream_with_the_real_transport(tmp_path, monkeypatch, encoding):
    from scripts import collect_factor_evidence as collector

    payload = b'{"prices":[10,11,12],"complete":true}'
    compressed = gzip.compress(payload) if encoding == "gzip" else zlib.compress(payload)
    observed = []

    class Handler(BaseHTTPRequestHandler):
        protocol_version = "HTTP/1.1"

        def do_GET(self):
            observed.append(dict(self.headers))
            self.send_response(200)
            self.send_header("Content-Encoding", encoding)
            self.send_header("Transfer-Encoding", "chunked")
            self.end_headers()
            for start in range(0, len(compressed), 7):
                chunk = compressed[start:start + 7]
                self.wfile.write(f"{len(chunk):x}\r\n".encode() + chunk + b"\r\n")
            self.wfile.write(b"0\r\n\r\n")
            self.wfile.flush()

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base = f"http://127.0.0.1:{server.server_port}/"
    monkeypatch.setattr(collector, "BASE", base)
    monkeypatch.setenv("NO_PROXY", "127.0.0.1")
    monkeypatch.setenv("no_proxy", "127.0.0.1")
    try:
        with requests.Session() as session:
            archive = collector.PublicArchive(tmp_path, session=session, timeout=2)
            raw, metadata = archive.fetch("sample", base + "prices")
            assert raw == payload
            assert metadata["status"] == 200
            assert metadata["body_bytes"] == len(payload)
            assert metadata["body_sha256"] == hashlib.sha256(payload).hexdigest()
            assert archive.fetch("sample", base + "prices") == (raw, metadata)
        assert len(observed) == 1
        assert not {key.lower() for key in observed[0]} & {"authorization", "proxy-authorization"}
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)
    assert not worker.is_alive()


def test_aiohttp_websocket_echo_and_normal_close_use_the_installed_parser():
    import aiohttp
    from aiohttp import web

    async def scenario():
        async def echo(request):
            websocket = web.WebSocketResponse()
            await websocket.prepare(request)
            await websocket.send_json({"status": "ready"})
            async for message in websocket:
                if message.type == aiohttp.WSMsgType.TEXT:
                    await websocket.send_str(message.data)
                    await websocket.close()
            return websocket

        app = web.Application()
        app.router.add_get("/stream", echo)
        runner = web.AppRunner(app)
        await runner.setup()
        import socket

        listener = socket.socket()
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
        try:
            await web.SockSite(runner, listener).start()
            async with aiohttp.ClientSession() as session:
                async with session.ws_connect(f"http://127.0.0.1:{port}/stream", heartbeat=1) as websocket:
                    assert await websocket.receive_json() == {"status": "ready"}
                    await websocket.send_str("complete-bar")
                    assert (await websocket.receive()).data == "complete-bar"
                    assert (await websocket.receive()).type == aiohttp.WSMsgType.CLOSE
        finally:
            await runner.cleanup()
            listener.close()

    asyncio.run(asyncio.wait_for(scenario(), timeout=10))


@pytest.mark.parametrize("kind", ["rsa", "ed25519-raw", "ed25519-pem"])
def test_ccxt_private_key_signatures_verify_with_the_installed_backend(kind):
    from ccxt.base.exchange import Exchange
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import ed25519, padding, rsa

    payload = b"symbol=BTCUSDT&side=BUY&timestamp=1"
    key = (rsa.generate_private_key(public_exponent=65537, key_size=2048) if kind == "rsa"
           else ed25519.Ed25519PrivateKey.generate())
    if kind == "ed25519-raw":
        secret = key.private_bytes(serialization.Encoding.Raw, serialization.PrivateFormat.Raw,
                                   serialization.NoEncryption())
    else:
        secret = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                   serialization.NoEncryption()).decode("ascii")
    signature = base64.b64decode(Exchange.rsa(payload.decode("ascii"), secret) if kind == "rsa"
                                 else Exchange.eddsa(payload, secret))
    if kind == "rsa":
        key.public_key().verify(signature, payload, padding.PKCS1v15(), hashes.SHA256())
    else:
        key.public_key().verify(signature, payload)


def test_ccxt_hmac_credentials_keep_the_same_signature():
    from ccxt.base.exchange import Exchange

    payload, secret = b"timestamp=1", b"synthetic-fixture-secret"
    assert Exchange.hmac(payload, secret, hashlib.sha256) == hmac.new(secret, payload, hashlib.sha256).hexdigest()


def test_generated_png_can_be_read_and_embedded_in_an_excel_report():
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as WorkbookImage
    from PIL import Image

    figure = Figure(figsize=(2, 1), dpi=50)
    FigureCanvasAgg(figure)
    figure.subplots().plot([0, 1, 2], [100, 102, 101])
    png = io.BytesIO()
    figure.savefig(png, format="png")
    with Image.open(io.BytesIO(png.getvalue())) as image:
        assert image.format == "PNG"
        assert image.size == (100, 50)
        image.load()
    workbook = Workbook()
    workbook.active.add_image(WorkbookImage(io.BytesIO(png.getvalue())), "A1")
    report = io.BytesIO()
    workbook.save(report)
    report.seek(0)
    restored = load_workbook(report)
    assert len(restored.active._images) == 1
    restored.close()


def test_provider_html_selection_preserves_result_order():
    from bs4 import BeautifulSoup

    document = BeautifulSoup('<table><tr class="quote"><td>BTC</td></tr>'
                             '<tr class="quote"><td>ETH</td></tr></table>', "html.parser")
    assert [cell.get_text() for cell in document.select("table tr.quote > td")] == ["BTC", "ETH"]


@pytest.mark.parametrize("entrypoint", ["domain", "label", "codec", "incremental_codec"])
def test_internationalized_domains_round_trip_through_supported_apis(entrypoint):
    import codecs

    import idna
    import idna.codec  # Register the third-party idna2008 codec.

    text, expected = "bücher.example", b"xn--bcher-kva.example"
    if entrypoint == "label":
        assert idna.alabel("bücher") == b"xn--bcher-kva"
        assert idna.ulabel(b"xn--bcher-kva") == "bücher"
    elif entrypoint == "codec":
        assert codecs.encode(text, "idna2008") == expected
        assert codecs.decode(expected, "idna2008") == text
    elif entrypoint == "incremental_codec":
        encoder = codecs.getincrementalencoder("idna2008")()
        decoder = codecs.getincrementaldecoder("idna2008")()
        assert encoder.encode("bücher.") + encoder.encode("example", final=True) == expected
        assert decoder.decode(b"xn--bcher-kva.") + decoder.decode(b"example", final=True) == text
    else:
        assert idna.encode(text) == expected
        assert idna.decode(expected) == text


@pytest.mark.parametrize("client", ["requests", "yarl"])
def test_http_clients_prepare_unicode_hosts_before_transport(client):
    from yarl import URL

    address = "https://BÜCHER.example/prices?symbol=BTC"
    prepared = (requests.Request("GET", address).prepare().url if client == "requests" else str(URL(address)))
    assert prepared == "https://xn--bcher-kva.example/prices?symbol=BTC"


@pytest.mark.parametrize("backend", ["multidict", "multidict._multidict_py"])
@pytest.mark.parametrize("container", ["MultiDict", "CIMultiDict"])
def test_header_containers_preserve_duplicates_and_copy_isolation(backend, container):
    module = importlib.import_module(backend)
    headers = getattr(module, container)([("X-Quote", "100"), ("X-Quote", "101")])
    assert headers.getall("X-Quote") == ["100", "101"]
    copied = headers.copy()
    copied.add("X-Quote", "102")
    assert copied.getall("X-Quote") == ["100", "101", "102"]
    assert headers.getall("X-Quote") == ["100", "101"]
    proxy = getattr(module, container + "Proxy")(headers)
    assert list(proxy.items()) == [("X-Quote", "100"), ("X-Quote", "101")]
    if container == "CIMultiDict":
        assert proxy.getall("x-quote") == ["100", "101"]


@pytest.mark.parametrize("selector,expected", [
    ("ul \t li", ["BTC", "ETH"]),
    ("ul > li /*comment*/", ["BTC", "ETH"]),
    (r'li[data-code="A\20 B"]', ["BTC"]),
    (r'li[data-code=A\20 B]', ["BTC"]),
    ("li:nth-child(2)", ["ETH"]),
])
def test_css_selector_variants_preserve_compiled_and_html_results(selector, expected):
    from bs4 import BeautifulSoup
    import soupsieve

    document = BeautifulSoup('<ul><li data-code="A B">BTC</li><li>ETH</li></ul>', "html.parser")
    assert [cell.get_text() for cell in document.select(selector)] == expected
    assert [cell.get_text() for cell in soupsieve.compile(selector).select(document)] == expected
